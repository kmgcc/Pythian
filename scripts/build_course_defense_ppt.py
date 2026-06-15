from __future__ import annotations

from pathlib import Path
import math

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIG_DIR = ROOT / "python课设_代码" / "outputs" / "figures"
RES_DIR = ROOT / "python课设_代码" / "outputs" / "results"
OUT = ROOT / "第24组_课程答辩.pptx"


PAPER = "FAFAF8"
INK = "0A0A0A"
GREY_1 = "F0F0EE"
GREY_2 = "D4D4D2"
GREY_3 = "737373"
ACCENT = "002FA7"
WHITE = "FFFFFF"
SUCCESS = "2F855A"
WARN = "B7791F"

FONT = "Microsoft YaHei"
FONT_LATIN = "Aptos"
FONT_MONO = "Consolas"


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.replace("#", "")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_run_font(run, size: float, color: str = INK, bold: bool = False, font: str = FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font: str = FONT,
    line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Cm(0)
    tf.margin_right = Cm(0)
    tf.margin_top = Cm(0)
    tf.margin_bottom = Cm(0)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    first = True
    for part in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = part
        set_run_font(run, size, color, bold, font)
    return box


def add_title(slide, title: str, index: str, kicker: str = "COURSE DEFENSE"):
    add_text(slide, f"{kicker} · {index}", 0.92, 0.24, 5.6, 0.30, 7.7, GREY_3, True, font=FONT_MONO)
    add_text(slide, title, 0.92, 0.73, 22.0, 0.76, 17.6, INK, True)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Cm(0.92), Cm(1.67), Cm(11.2), Cm(0.035))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(ACCENT)
    line.line.fill.background()


def add_footer(slide, idx: int, text: str = "基于真实天光数据的自然光光谱估计与室内照明补偿设计"):
    add_text(slide, text, 0.92, 18.42, 10.5, 0.35, 7.2, GREY_3, False, font=FONT_MONO)
    add_text(slide, f"{idx:02d}/10", 30.15, 18.42, 2.0, 0.35, 7.2, GREY_3, False, align=PP_ALIGN.RIGHT, font=FONT_MONO)


def set_bg(slide, color: str = PAPER):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(color)


def add_rect(slide, x, y, w, h, fill=GREY_1, line=None, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Cm(x), Cm(y), Cm(w), Cm(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.6)
    else:
        shp.line.fill.background()
    return shp


def add_chip(slide, text, x, y, w, fill=GREY_1, color=INK):
    shp = add_rect(slide, x, y, w, 0.62, fill=fill, line=None)
    add_text(slide, text, x + 0.18, y + 0.16, w - 0.36, 0.32, 8.5, color, True, align=PP_ALIGN.CENTER)
    return shp


def add_bullets(slide, items, x, y, w, row_h=0.82, size=11.5, color=INK, accent=ACCENT):
    for i, item in enumerate(items):
        yy = y + i * row_h
        sq = add_rect(slide, x, yy + 0.16, 0.13, 0.13, fill=accent)
        sq.line.fill.background()
        add_text(slide, item, x + 0.34, yy, w - 0.34, row_h, size, color, False, line_spacing=1.07)


def add_caption(slide, text, x, y, w):
    add_text(slide, text, x, y, w, 0.38, 7.2, GREY_3, False, font=FONT_MONO)


def add_picture_fit(slide, path: Path, x, y, w, h, border: bool = False):
    with Image.open(path) as im:
        iw, ih = im.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio >= box_ratio:
        width = w
        height = w / img_ratio
        xx = x
        yy = y + (h - height) / 2
    else:
        height = h
        width = h * img_ratio
        xx = x + (w - width) / 2
        yy = y
    pic = slide.shapes.add_picture(str(path), Cm(xx), Cm(yy), Cm(width), Cm(height))
    if border:
        frame = add_rect(slide, x, y, w, h, fill=PAPER, line=GREY_2)
        slide.shapes._spTree.remove(frame._element)
        slide.shapes._spTree.insert(2, frame._element)
    return pic


def add_metric(slide, label, value, x, y, w=3.0, color=ACCENT):
    add_text(slide, value, x, y, w, 0.64, 20, color, True, font=FONT_LATIN)
    add_text(slide, label, x, y + 0.76, w, 0.36, 7.6, GREY_3, True, font=FONT_MONO)


def add_card(slide, title, body, x, y, w, h, fill=GREY_1, title_color=ACCENT):
    add_rect(slide, x, y, w, h, fill=fill, line=None)
    add_text(slide, title, x + 0.35, y + 0.30, w - 0.70, 0.38, 10.2, title_color, True)
    add_text(slide, body, x + 0.35, y + 0.92, w - 0.70, h - 1.10, 9.6, INK, False, line_spacing=1.05)


def add_takeaway(slide, text, x, y, w, h=1.0, fill=ACCENT, color=WHITE):
    add_rect(slide, x, y, w, h, fill=fill)
    add_text(slide, text, x + 0.35, y + 0.23, w - 0.70, h - 0.34, 10.2, color, True, line_spacing=1.0)


def add_explain_box(slide, title, body, x, y, w, h, fill=GREY_1, title_color=ACCENT):
    add_rect(slide, x, y, w, h, fill=fill)
    add_text(slide, title, x + 0.35, y + 0.25, w - 0.70, 0.34, 9.4, title_color, True)
    add_text(slide, body, x + 0.35, y + 0.78, w - 0.70, h - 0.92, 8.7, INK, False, line_spacing=1.03)


def add_arrow(slide, x1, y1, x2, y2, color=ACCENT, width=1.2):
    con = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    con.line.color.rgb = rgb(color)
    con.line.width = Pt(width)
    con.line.end_arrowhead = True
    return con


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_rect(slide, 0, 0, 3.15, 19.05, fill=ACCENT)
    add_text(slide, "PYTHON\nCOURSE\nDESIGN", 0.62, 0.70, 2.0, 2.2, 9.0, WHITE, True, font=FONT_MONO)
    add_text(slide, "第24组", 0.62, 17.62, 1.9, 0.4, 9.0, WHITE, True, font=FONT_MONO)
    add_text(slide, "基于真实天光数据的\n自然光光谱估计与\n室内照明补偿设计", 4.15, 2.08, 16.5, 3.95, 28, INK, True, line_spacing=0.90)
    add_text(slide, "课程名称：Python 应用开发基础\n小组成员：李安逸、黄奕滔", 4.22, 6.70, 10.5, 1.25, 12, INK, False, line_spacing=1.2)
    add_text(slide, "内容主线：光谱问题 → 数据建模 → LED 补偿应用", 4.22, 7.92, 15.0, 0.45, 9.7, GREY_3, True)
    add_rect(slide, 4.20, 8.70, 7.8, 2.15, fill=GREY_1)
    add_metric(slide, "VISIBLE WAVELENGTH POINTS", "41", 4.62, 9.03, 2.3)
    add_metric(slide, "REAL SAMPLES", "5664", 7.06, 9.03, 2.6)
    add_metric(slide, "PCA CUM. VAR.", "98.3%", 10.02, 9.03, 2.0)
    add_explain_box(
        slide,
        "内容概览",
        "本项目先说明为什么亮度和色温不足以描述照明质量，再说明如何用真实数据训练模型，最后展示预测光谱怎样进入 LED 补偿应用。",
        4.20,
        11.35,
        11.25,
        2.35,
    )
    add_explain_box(
        slide,
        "项目定位",
        "本项目是课程设计中的算法原型，重点展示 Python 数据处理、PCA 降维、模型对比和可视化结果，不把它夸大为已经落地的硬件系统。",
        4.20,
        14.10,
        11.25,
        2.25,
        fill=PAPER,
    )
    add_picture_fit(slide, FIG_DIR / "base_spectrum.png", 16.0, 8.58, 14.6, 7.4)
    add_caption(slide, "图：实测自然光基准光谱曲线", 16.0, 16.18, 8.0)
    add_text(slide, "答辩时间约 6 分钟", 25.1, 0.72, 5.0, 0.4, 8.5, GREY_3, True, align=PP_ALIGN.RIGHT, font=FONT_MONO)
    return slide


def slide_background(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "研究背景：照明质量不能只看亮度和色温", "02")
    add_bullets(
        slide,
        [
            "人长期处于室内，学习、办公和居家场景都依赖人工照明。",
            "照度看亮度，色温看冷暖；不能描述完整波长分布。",
            "光谱表示各波长能量分布，更接近光照质量本身。",
            "同亮度、同色温下，光谱形状仍可能明显不同。",
        ],
        1.05,
        2.60,
        10.8,
        row_h=0.84,
        size=10.8,
    )
    add_rect(slide, 1.05, 6.18, 10.65, 2.45, fill=ACCENT)
    add_text(slide, "读图方式：不同天气曲线不完全重合，\n说明自然光光谱会随环境变化。\n因此要用“光谱”而不只用亮度/色温描述。", 1.45, 6.50, 9.65, 1.55, 11.4, WHITE, True, line_spacing=1.0)
    add_explain_box(
        slide,
        "研究切入点",
        "研究从室内照明场景出发：人长期处于室内，照明质量会影响学习、办公和视觉舒适。亮度和色温虽然直观，但无法说明每个波长的能量分布。",
        1.05,
        9.18,
        10.65,
        3.10,
        fill=GREY_1,
    )
    add_explain_box(
        slide,
        "建模问题",
        "天气变化会改变自然光光谱形状，因此可以进一步尝试利用天气、时间和太阳位置等环境特征去估计这条光谱曲线。",
        1.05,
        12.62,
        10.65,
        2.25,
        fill=PAPER,
    )
    add_picture_fit(slide, FIG_DIR / "weather_spectrum_compare.png", 13.05, 2.14, 18.0, 13.15)
    add_caption(slide, "关键图 1：不同天气下平均相对光谱存在差异", 13.05, 15.55, 13.5)
    add_footer(slide, 2)


def slide_goal(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "项目目标：由环境特征估计自然光相对光谱", "03")
    add_card(slide, "输入", "不直接依赖光谱仪\n天气类别、云量、湿度、温度、降水\n时间、月份、地点\n太阳高度角/方位角、室外照度", 1.0, 3.1, 8.2, 5.8)
    add_card(slide, "输出", "380–780 nm 可见光范围\n每 10 nm 一个点，共 41 维\n关注相对光谱形状\n先预测 PCA 系数，再还原曲线", 12.25, 3.1, 8.2, 5.8)
    add_card(slide, "应用", "计算“目标光谱 - 当前自然光”\n输出七通道 LED 推荐比例\n对比多通道与传统双色温\n形成预测、补偿、评价闭环", 23.5, 3.1, 8.2, 5.8)
    add_arrow(slide, 9.52, 5.95, 11.55, 5.95)
    add_arrow(slide, 20.78, 5.95, 22.78, 5.95)
    add_explain_box(
        slide,
        "设计思路",
        "输入端尽量选择低成本、容易获得的环境特征；输出端不直接给一个色温值，而是还原完整相对光谱。这样预测结果才能继续用于光谱级补偿。",
        3.7,
        9.05,
        24.6,
        1.25,
        fill=GREY_1,
    )
    add_rect(slide, 3.7, 10.8, 24.6, 3.15, fill=PAPER, line=GREY_2)
    labels = ["环境特征", "机器学习回归", "PCA逆变换", "光谱曲线", "LED补偿"]
    xs = [4.45, 9.35, 14.45, 19.25, 24.1]
    for i, (label, x) in enumerate(zip(labels, xs)):
        fill = ACCENT if i in [1, 4] else GREY_1
        color = WHITE if fill == ACCENT else INK
        add_chip(slide, label, x, 11.96, 3.1, fill=fill, color=color)
        if i < len(labels) - 1:
            add_arrow(slide, x + 3.28, 12.27, xs[i + 1] - 0.15, 12.27, color=GREY_3, width=0.75)
    add_takeaway(slide, "一句话目标：用低成本环境特征估计自然光相对光谱，并服务于室内照明补偿。", 3.7, 14.45, 24.6, h=0.85)
    add_footer(slide, 3)


def slide_data(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "数据来源：公开真实数据构建完整样本表", "04")
    add_card(slide, "SKYSPECTRA 实测天光光谱", "Zenodo record 8147546\n水平天光光谱 + 站点、时间、天空状况、太阳位置元数据\n重采样到 380–780 nm / 10 nm", 1.0, 2.45, 9.45, 5.35, fill=GREY_1)
    add_card(slide, "Open-Meteo 历史天气 API", "按观测站经纬度和日期范围查询\n云量、湿度、温度、降水\n按最近整点与光谱测量对齐", 11.3, 2.45, 9.45, 5.35, fill=GREY_1)
    add_card(slide, "公开 LED 光谱数据", "用于构造七通道 LED 补偿模型\n深蓝、青、绿、琥珀、红、暖白、冷白\n只用于应用补偿，不参与模型训练", 21.6, 2.45, 9.45, 5.35, fill=GREY_1)
    add_explain_box(
        slide,
        "数据处理后的建模表",
        "最终每一行样本都包含：观测时间和地点、天气特征、太阳位置、室外照度，以及 41 个波长点的相对光谱强度。也就是说，模型看到的是“环境条件 → 光谱形状”的配对样本。",
        1.0,
        8.10,
        30.05,
        1.08,
        fill=GREY_1,
    )
    add_rect(slide, 1.0, 9.45, 30.05, 4.3, fill=PAPER, line=GREY_2)
    add_metric(slide, "样本量", "5664", 2.0, 10.25, 4.0)
    add_metric(slide, "光谱维度", "41", 7.05, 10.25, 3.8)
    add_metric(slide, "时间范围", "2016–2018", 11.25, 10.25, 5.8)
    add_metric(slide, "观测站", "2", 18.25, 10.25, 2.8)
    add_metric(slide, "主要站点", "法国沃昂夫兰", 22.0, 10.25, 5.9)
    add_text(slide, "使用原则：估计模型只用天光光谱和环境特征训练；LED 光谱只在补偿应用中使用，避免数据角色混淆。", 2.05, 12.82, 27.2, 0.45, 9.3, INK, True)
    add_text(slide, "边界说明：天气数据为按地点和时间对齐的历史天气特征，不等同于现场同步气象测量。", 2.05, 13.42, 27.2, 0.45, 9.3, GREY_3, False)
    add_footer(slide, 4)


def slide_flow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "总体技术流程：从真实数据到照明补偿闭环", "05")
    steps = [
        ("01", "数据下载", "SKYSPECTRA/API"),
        ("02", "清洗重采样", "统一波长网格"),
        ("03", "天气对齐", "最近整点匹配"),
        ("04", "特征工程", "标准化/独热编码"),
        ("05", "PCA", "41 维降到 5 维"),
        ("06", "五模型训练", "统一指标对比"),
        ("07", "光谱还原", "inverse transform"),
        ("08", "LED 补偿", "七通道比例"),
    ]
    x0, y0, w, h, gap = 1.05, 3.20, 6.65, 2.25, 1.0
    for i, (num, label, desc) in enumerate(steps):
        row = 0 if i < 4 else 1
        col = i if i < 4 else 7 - i
        x = x0 + col * (w + gap)
        y = y0 + row * 4.45
        fill = ACCENT if i in [4, 7] else GREY_1
        txt_color = WHITE if fill == ACCENT else INK
        add_rect(slide, x, y, w, h, fill=fill)
        add_text(slide, num, x + 0.35, y + 0.30, 1.0, 0.4, 9.4, txt_color, True, font=FONT_MONO)
        add_text(slide, label, x + 0.35, y + 0.92, w - 0.7, 0.42, 12.4, txt_color, True)
        add_text(slide, desc, x + 0.35, y + 1.48, w - 0.7, 0.34, 8.1, txt_color if fill == ACCENT else GREY_3, False, font=FONT_MONO)
        if i < 3:
            add_arrow(slide, x + w + 0.15, y + 1.12, x + w + gap - 0.20, y + 1.12, color=GREY_3, width=0.75)
        if i == 3:
            add_arrow(slide, x + w / 2, y + h + 0.2, x + w / 2, y + 4.05, color=GREY_3, width=0.75)
        if 4 <= i < 7:
            x_next = x0 + (col - 1) * (w + gap)
            add_arrow(slide, x - 0.18, y + 1.12, x_next + w + 0.20, y + 1.12, color=GREY_3, width=0.75)
    add_rect(slide, 1.05, 13.40, 30.0, 1.75, fill=PAPER, line=GREY_2)
    add_text(slide, "核心逻辑：真实数据支撑建模，PCA 降低输出维度，模型预测的光谱再进入七通道 LED 补偿。", 1.55, 13.66, 28.7, 0.42, 11.3, INK, True)
    add_text(slide, "流程特点：每一步都有对应的数据、模型或图表输出，能够支撑完整的课程演示。", 1.55, 14.33, 28.7, 0.38, 9.3, GREY_3, False)
    add_explain_box(
        slide,
        "流程关系",
        "前半部分把真实数据整理成训练集，中间用 PCA 和回归模型完成光谱估计，后半部分把预测结果接到照明补偿应用里。",
        1.05,
        10.65,
        30.0,
        1.90,
        fill=GREY_1,
    )
    add_footer(slide, 5)


def slide_pca(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "PCA 与特征工程：把 41 维光谱压缩成可学习目标", "06")
    add_picture_fit(slide, FIG_DIR / "pca_variance.png", 1.0, 2.25, 14.0, 9.3)
    add_picture_fit(slide, FIG_DIR / "feature_importance.png", 16.7, 2.25, 14.2, 9.3)
    add_caption(slide, "PCA 解释方差：前 5 个主成分累计约 98.3%", 1.0, 11.78, 12.8)
    add_caption(slide, "随机森林特征重要性：室外照度、云量、太阳位置影响明显", 16.7, 11.78, 14.0)
    add_explain_box(
        slide,
        "为什么要 PCA",
        "光谱曲线相邻波长之间变化连续、相关性强，直接预测 41 个输出值会增加模型难度。PCA 把主要变化压缩成少量系数，既保留形状信息，也便于回归模型学习。",
        1.0,
        12.20,
        14.0,
        1.72,
        fill=PAPER,
    )
    add_explain_box(
        slide,
        "特征怎么理解",
        "特征重要性不是物理因果结论，只说明在当前数据和随机森林模型中，室外照度、云量、太阳方位角等信息对预测光谱形状更有帮助。",
        16.7,
        12.20,
        14.2,
        1.72,
        fill=PAPER,
    )
    add_rect(slide, 1.0, 14.25, 30.0, 1.35, fill=GREY_1)
    add_bullets(
        slide,
        [
            "输入特征：天气、时间、地点、太阳位置、室外照度；类别特征独热编码，数值特征标准化。",
            "输出目标：模型不直接预测 41 维光谱，而是预测 PCA 主成分系数，再逆变换还原。",
            "选择依据：前 5 个主成分累计解释约 98.3% 方差，兼顾信息保留和训练难度。",
        ],
        1.45,
        14.45,
        28.3,
        row_h=0.38,
        size=8.4,
    )
    add_footer(slide, 6)


def slide_models(prs, metrics: pd.DataFrame):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "模型训练与对比：五个回归模型在同一数据上评价", "07")
    add_picture_fit(slide, FIG_DIR / "model_compare.png", 1.0, 2.15, 19.5, 7.3)
    add_caption(slide, "指标柱状图：RMSE、R²、训练时间", 1.0, 9.55, 10.0)
    add_rect(slide, 1.0, 10.25, 19.5, 2.0, fill=GREY_1)
    add_bullets(
        slide,
        [
            "MAE 表示平均偏差；RMSE 对较大误差更敏感；R² 反映解释能力。",
            "五个模型输入和输出完全一致，因此比较的是模型本身的拟合能力。",
            "选择模型时不只看 R²，也要结合 RMSE、训练成本和预测稳定性。",
        ],
        1.45,
        10.68,
        18.5,
        row_h=0.47,
        size=8.6,
    )
    table_data = [["模型", "MAE", "RMSE", "R²", "训练/s"]]
    order = ["Random Forest", "KNN", "MLP", "Decision Tree", "Linear Regression"]
    m = metrics.set_index("model")
    for name in order:
        row = m.loc[name]
        table_data.append([name, f"{row['MAE']:.4f}", f"{row['RMSE']:.4f}", f"{row['R2']:.3f}", f"{row['train_time_s']:.2f}"])
    rows, cols = len(table_data), len(table_data[0])
    table_shape = slide.shapes.add_table(rows, cols, Cm(21.25), Cm(2.35), Cm(9.85), Cm(5.7))
    table = table_shape.table
    widths = [3.45, 1.45, 1.55, 1.35, 1.55]
    for i, width in enumerate(widths):
        table.columns[i].width = Cm(width)
    for r, row in enumerate(table_data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            cell.margin_left = Cm(0.08)
            cell.margin_right = Cm(0.08)
            cell.margin_top = Cm(0.06)
            cell.margin_bottom = Cm(0.06)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(ACCENT if r == 0 else (GREY_1 if r % 2 else PAPER))
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
                for run in p.runs:
                    set_run_font(run, 7.6 if r else 7.8, WHITE if r == 0 else INK, bold=(r == 0 or (r == 1 and c == 0)), font=FONT_LATIN if c == 0 else FONT_MONO)
    add_rect(slide, 21.25, 8.65, 9.85, 2.35, fill=ACCENT)
    add_text(slide, "最终按 RMSE 最小选择随机森林", 21.75, 9.05, 8.8, 0.48, 12.6, WHITE, True)
    add_text(slide, "随机森林 RMSE 最小，综合稳定性更好。\nKNN 的 MAE 略低，但 RMSE 与预测耗时不占优；MLP 的 R² 略高，但误差更大。", 21.75, 9.68, 8.75, 1.05, 8.1, WHITE, False, line_spacing=1.0)
    add_explain_box(
        slide,
        "模型选择结论",
        "线性回归明显落后，说明环境特征到光谱形状不是简单线性关系。单棵决策树也不如随机森林，体现了集成模型在稳定性上的优势。",
        1.0,
        12.70,
        30.1,
        1.55,
        fill=PAPER,
    )
    add_footer(slide, 7)


def slide_prediction(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "光谱预测结果：环境特征可以还原相对光谱形状", "08")
    add_picture_fit(slide, FIG_DIR / "prediction_compare.png", 1.0, 2.15, 21.2, 12.2)
    add_rect(slide, 23.1, 2.28, 8.0, 5.35, fill=GREY_1)
    add_text(slide, "预测路径", 23.55, 2.72, 7.1, 0.35, 10.2, ACCENT, True)
    add_bullets(
        slide,
        [
            "输入环境特征",
            "预测 5 个 PCA 系数",
            "PCA 逆变换还原 41 维光谱",
            "与实测光谱曲线比较",
        ],
        23.55,
        3.42,
        6.95,
        row_h=0.70,
        size=9.4,
    )
    add_rect(slide, 23.1, 8.20, 8.0, 3.4, fill=ACCENT)
    add_text(slide, "结果观察", 23.55, 8.62, 7.1, 0.35, 10.2, WHITE, True)
    add_text(slide, "示例样本中，预测曲线与实测曲线基本重合。\n模型能够学习自然光相对光谱的主要变化趋势。", 23.55, 9.30, 7.0, 1.32, 8.6, WHITE, False, line_spacing=1.0)
    add_rect(slide, 23.1, 12.18, 8.0, 1.42, fill=GREY_1)
    add_text(slide, "结果边界：预测的是相对光谱形状，不能替代现场光谱仪的绝对测量。", 23.55, 12.50, 7.1, 0.55, 8.5, INK, True, line_spacing=1.0)
    add_explain_box(
        slide,
        "结果解读",
        "图中两条线越接近，说明还原后的光谱形状越接近实测结果。这里不是证明模型永远准确，而是说明在测试样本上，它能抓住自然光光谱的主要起伏。",
        1.0,
        14.95,
        21.2,
        1.25,
        fill=GREY_1,
    )
    add_caption(slide, "关键图 2：实测光谱 vs 模型预测光谱", 1.0, 14.58, 13.8)
    add_footer(slide, 8)


def slide_compensation(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "照明补偿结果：七通道 LED 比双色温更细致", "09")
    add_picture_fit(slide, FIG_DIR / "application_spectra_compare.png", 1.0, 2.10, 20.7, 12.15)
    add_caption(slide, "关键图 3：目标光谱、预测自然光、补偿后光谱与传统双色温方案对比", 1.0, 14.45, 18.0)
    add_rect(slide, 22.4, 2.22, 8.75, 4.7, fill=GREY_1)
    add_text(slide, "补偿思路", 22.85, 2.62, 7.8, 0.35, 10.2, ACCENT, True)
    add_text(slide, "目标光谱 - 当前自然光贡献\n= 需要人工补偿的光谱", 22.85, 3.12, 7.7, 0.98, 10.3, INK, True, line_spacing=1.0)
    add_text(slide, "图表解读：黑线目标、蓝虚线自然光、\n绿线多通道补偿、橙线传统双色温。", 22.85, 4.42, 7.7, 0.95, 8.1, INK, False, line_spacing=1.0)
    add_text(slide, "算法：非负最小二乘求 7 通道比例。", 22.85, 5.82, 7.7, 0.38, 8.3, GREY_3, False)
    add_rect(slide, 22.4, 7.55, 8.75, 4.95, fill=ACCENT)
    add_text(slide, "七通道 LED", 22.85, 7.98, 7.8, 0.35, 10.2, WHITE, True)
    add_text(slide, "深蓝/蓝光、青光、绿光、琥珀光、红光、暖白、冷白\n\n通道更多，可按波段局部补偿；\n双色温主要调冷暖，光谱自由度较低。", 22.85, 8.60, 7.6, 2.55, 8.4, WHITE, False, line_spacing=1.0)
    add_text(slide, "课程定位：这里展示的是算法设计方案，暂未接入真实灯具硬件。", 22.85, 12.85, 7.6, 0.52, 8.4, GREY_3, True)
    add_explain_box(
        slide,
        "补偿结论",
        "多通道 LED 的优势在于自由度更多：不同波段可以分别补。传统双色温方案主要改变冷暖比例，因此在某些波段会更难贴近目标光谱。",
        1.0,
        14.95,
        20.7,
        1.25,
        fill=GREY_1,
    )
    add_footer(slide, 9)


def slide_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "项目特色、不足与分工", "10")
    add_card(
        slide,
        "项目特色",
        "真实公开数据\nPCA 降维表示光谱\n五模型统一对比\n预测结果进入 LED 补偿闭环",
        1.0,
        2.35,
        9.7,
        5.2,
        fill=GREY_1,
    )
    add_card(
        slide,
        "主要不足",
        "样本集中在少数站点\n天气数据非现场同步\n模型不能替代现场光谱仪\n未接入真实硬件",
        11.45,
        2.35,
        9.7,
        5.2,
        fill=GREY_1,
        title_color=WARN,
    )
    add_card(
        slide,
        "后续改进",
        "补充更多地区实测数据\n接入实时天气与廉价传感器\n结合智能灯具进行硬件验证\n进一步评估舒适度和显色效果",
        21.9,
        2.35,
        9.2,
        5.2,
        fill=GREY_1,
        title_color=SUCCESS,
    )
    add_rect(slide, 1.0, 9.55, 30.1, 4.7, fill=PAPER, line=GREY_2)
    add_text(slide, "小组分工", 1.55, 10.05, 5.0, 0.45, 12, ACCENT, True)
    add_text(slide, "黄奕滔", 2.0, 11.05, 3.0, 0.42, 12, INK, True)
    add_text(slide, "数据处理、模型训练、天气 API 获取、Streamlit 页面、代码注释与运行结果整理", 5.2, 11.03, 24.5, 0.5, 10.2, INK, False)
    add_text(slide, "李安逸", 2.0, 12.22, 3.0, 0.42, 12, INK, True)
    add_text(slide, "课程报告、背景与方法说明、PCA 与模型分析、PPT 逻辑整理与答辩表达", 5.2, 12.20, 24.5, 0.5, 10.2, INK, False)
    add_text(slide, "共同完成：课程答辩、PPT、最终材料整合。", 5.2, 13.36, 24.5, 0.5, 10.2, GREY_3, False)
    add_explain_box(
        slide,
        "项目总结",
        "本项目是一个完整但谨慎的课程项目：数据是真实的，流程是可复现的，结果能进入应用演示；同时也承认样本、同步天气和硬件验证上的不足。",
        1.0,
        14.20,
        30.1,
        1.05,
        fill=GREY_1,
    )
    add_rect(slide, 1.0, 15.45, 30.1, 1.45, fill=ACCENT)
    add_text(slide, "总结：本项目把真实天光数据、机器学习光谱估计和室内照明补偿连成了一条可复现的课程设计流程；它是课程层面的算法原型，不夸大为硬件成品。", 1.55, 15.80, 29.0, 0.68, 11.1, WHITE, True, line_spacing=1.0)
    add_footer(slide, 10)


def build():
    metrics = pd.read_csv(RES_DIR / "model_metrics.csv")
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    # Drop default empty slide if present in some templates.
    while len(prs.slides) > 0:
        r_id = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]

    slide_cover(prs)
    slide_background(prs)
    slide_goal(prs)
    slide_data(prs)
    slide_flow(prs)
    slide_pca(prs)
    slide_models(prs, metrics)
    slide_prediction(prs)
    slide_compensation(prs)
    slide_closing(prs)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
